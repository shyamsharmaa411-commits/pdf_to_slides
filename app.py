import io
import json
import os
import streamlit as st
from pypdf import PdfReader
import google.generativeai as genai
from pptx import Presentation
from pptx.util import Inches

st.set_page_config(page_title="PDF to MCQ Slides Generator", layout="wide")
st.title("PDF to MCQ Slides Generator")

MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-3.6-flash")

def extract_pdf_text(uploaded_file):
    reader = PdfReader(uploaded_file)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def generate_mcqs(pdf_text, api_key):
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(MODEL_NAME)

    prompt = f"""
    Read the following PDF text and generate exactly 5 multiple-choice questions.
    Return only valid JSON in this exact format:
    [
      {{
        "question": "Question text",
        "options": ["A", "B", "C", "D"],
        "answer": "Correct option text"
      }}
    ]

    Rules:
    - Questions must be based only on the PDF content.
    - Each question must have 4 options.
    - Only one option is correct.
    - Return valid JSON only. No markdown. No extra text.

    PDF TEXT:
    {pdf_text[:12000]}
    """

    response = model.generate_content(prompt)
    result = response.text.strip()

    if result.startswith("```"):
        result = result.replace("```json", "").replace("```", "").strip()

    return json.loads(result)

def create_pptx(mcqs):
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "MCQ Quiz"
    title_slide.placeholders[1].text = "Generated from PDF"

    for i, mcq in enumerate(mcqs, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(11.5), Inches(1.2))
        title_box.text_frame.text = f"Q{i}: {mcq['question']}"
        title_box.text_frame.word_wrap = True

        y = 1.8
        for option in mcq["options"]:
            box = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.0), Inches(0.7))
            tf = box.text_frame
            tf.text = option
            tf.word_wrap = True
            y += 0.8

        answer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(10.0), Inches(0.8))
        answer_box.text_frame.text = f"Answer: {mcq['answer']}"

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()

uploaded_file = st.file_uploader("Choose a PDF file", type=["pdf"])

if uploaded_file is not None:
    st.success(f"PDF uploaded successfully: {uploaded_file.name}")

    extracted_text = extract_pdf_text(uploaded_file)

    with st.expander("View extracted PDF text"):
        st.write(extracted_text[:5000])

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        api_key = st.text_input("Enter Gemini API Key", type="password")

    if st.button("Generate MCQs & PPT"):
        if not api_key:
            st.warning("Please enter your Gemini API key.")
        else:
            try:
                mcqs = generate_mcqs(extracted_text, api_key)

                st.subheader("Generated MCQs")
                for i, mcq in enumerate(mcqs, start=1):
                    st.write(f"Q{i}. {mcq['question']}")
                    for option in mcq["options"]:
                        st.write(f"- {option}")
                    st.write(f"Correct Answer: {mcq['answer']}")
                    st.write("")

                ppt_bytes = create_pptx(mcqs)
                st.success("PowerPoint file generated successfully.")

                st.download_button(
                    label="Download PPT",
                    data=ppt_bytes,
                    file_name="mcq_slides.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )

            except Exception as e:
                st.error(f"Error: {e}")

else:
    st.info("Please upload a PDF file first.")